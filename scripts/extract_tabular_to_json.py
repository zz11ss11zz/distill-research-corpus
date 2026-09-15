from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook


def io_path(path: Path) -> Path:
    absolute = os.path.abspath(path)
    if os.name != "nt" or absolute.startswith("\\\\?\\"):
        return Path(absolute)
    if absolute.startswith("\\\\"):
        return Path("\\\\?\\UNC\\" + absolute[2:])
    return Path("\\\\?\\" + absolute)


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def dump_line(stream, value: Any) -> None:
    json.dump(value, stream, ensure_ascii=False, separators=(",", ":"))


def csv_reader(path: Path) -> tuple[str, csv.Dialect, Any, Iterator[list[str]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        stream = None
        try:
            stream = path.open("r", encoding=encoding, newline="", errors="strict")
            sample = stream.read(65536)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            except csv.Error:
                dialect = csv.excel
            return encoding, dialect, stream, csv.reader(stream, dialect)
        except (UnicodeError, OSError, csv.Error) as exc:
            last_error = exc
            if stream is not None:
                stream.close()
    raise RuntimeError(f"Unable to read CSV: {last_error}")


def write_csv_json(source: dict[str, Any], source_path: Path, output_path: Path) -> dict[str, Any]:
    encoding, dialect, input_stream, reader = csv_reader(source_path)
    row_count = 0
    cell_count = 0
    try:
        headers = next(reader, [])
        with output_path.open("w", encoding="utf-8", newline="\n") as stream:
            stream.write('{"schema_version":"1.0.0","source_id":')
            dump_line(stream, source["source_id"])
            stream.write(',"source_path":')
            dump_line(stream, source["path"])
            stream.write(',"sha256":')
            dump_line(stream, source["sha256"])
            stream.write(',"format":"csv","encoding":')
            dump_line(stream, encoding)
            stream.write(',"delimiter":')
            dump_line(stream, dialect.delimiter)
            stream.write(',"tables":[{"name":"data","headers":')
            dump_line(stream, headers)
            stream.write(',"rows":[')
            first = True
            for row in reader:
                if not first:
                    stream.write(",")
                dump_line(stream, [json_safe(value) for value in row])
                first = False
                row_count += 1
                cell_count += len(row)
            stream.write("]}]}")
    finally:
        input_stream.close()
    return {
        "source_id": source["source_id"],
        "source_path": source["path"],
        "sha256": source["sha256"],
        "format": "csv",
        "json_path": output_path.as_posix(),
        "table_count": 1,
        "row_count": row_count,
        "cell_count": cell_count,
        "headers": [{"table": "data", "columns": headers}],
    }


def write_xlsx_json(source: dict[str, Any], source_path: Path, output_path: Path) -> dict[str, Any]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    total_rows = 0
    total_cells = 0
    header_manifest = []
    try:
        with output_path.open("w", encoding="utf-8", newline="\n") as stream:
            stream.write('{"schema_version":"1.0.0","source_id":')
            dump_line(stream, source["source_id"])
            stream.write(',"source_path":')
            dump_line(stream, source["path"])
            stream.write(',"sha256":')
            dump_line(stream, source["sha256"])
            stream.write(',"format":"xlsx","tables":[')
            first_sheet = True
            for sheet in workbook.worksheets:
                if not first_sheet:
                    stream.write(",")
                first_sheet = False
                iterator = sheet.iter_rows(values_only=True)
                raw_headers = next(iterator, ())
                headers = [json_safe(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(raw_headers)]
                header_manifest.append({"table": sheet.title, "columns": headers})
                stream.write('{"name":')
                dump_line(stream, sheet.title)
                stream.write(',"headers":')
                dump_line(stream, headers)
                stream.write(',"rows":[')
                first_row = True
                for row in iterator:
                    if not first_row:
                        stream.write(",")
                    values = [json_safe(value) for value in row]
                    dump_line(stream, values)
                    first_row = False
                    total_rows += 1
                    total_cells += len(values)
                stream.write("]}")
            stream.write("]}")
    finally:
        workbook.close()
    return {
        "source_id": source["source_id"],
        "source_path": source["path"],
        "sha256": source["sha256"],
        "format": "xlsx",
        "json_path": output_path.as_posix(),
        "table_count": len(header_manifest),
        "row_count": total_rows,
        "cell_count": total_cells,
        "headers": header_manifest,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert every unique canonical CSV/XLSX cell into JSON.")
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--knowledge-dir", type=Path, default=Path("outputs/json/knowledge_distillation"))
    args = parser.parse_args()
    root = args.root.resolve()
    knowledge_dir = root / args.knowledge_dir
    inventory = json.loads((knowledge_dir / "source_inventory.json").read_text(encoding="utf-8"))["files"]
    duplicate_entries = json.loads((knowledge_dir / "deduplication.json").read_text(encoding="utf-8"))["duplicate_or_unreadable_entries"]
    duplicate_paths = {item["duplicate_path"] for item in duplicate_entries if item.get("reason") == "byte_identical_sha256"}

    seen_hashes: set[str] = set()
    sources = []
    for item in inventory:
        if item["category"] != "dataset" or item["extension"] not in {".csv", ".xlsx"}:
            continue
        if item["path"] in duplicate_paths or item["sha256"] in seen_hashes:
            continue
        seen_hashes.add(item["sha256"])
        source_id = "src_" + hashlib.sha256(item["sha256"].encode("utf-8")).hexdigest()[:20]
        sources.append({**item, "source_id": source_id})

    output_dir = knowledge_dir / "tabular"
    output_dir.mkdir(parents=True, exist_ok=True)
    prior_by_hash: dict[str, dict[str, Any]] = {}
    prior_manifest_path = knowledge_dir / "tabular_data_manifest.json"
    if prior_manifest_path.exists():
        try:
            prior_manifest = json.loads(prior_manifest_path.read_text(encoding="utf-8"))
            prior_by_hash = {item["sha256"]: item for item in prior_manifest.get("records", [])}
        except (json.JSONDecodeError, KeyError, TypeError):
            prior_by_hash = {}
    records = []
    failures = []
    reused = 0
    for index, source in enumerate(sources, 1):
        source_path = io_path(root / Path(source["path"]))
        output_path = output_dir / f"{source['source_id']}.json"
        try:
            prior = prior_by_hash.get(source["sha256"])
            if prior and output_path.exists():
                record = dict(prior)
                record.update({
                    "source_id": source["source_id"],
                    "source_path": source["path"],
                    "sha256": source["sha256"],
                    "json_path": output_path.relative_to(root).as_posix(),
                })
                reused += 1
            elif source["extension"] == ".csv":
                record = write_csv_json(source, source_path, output_path)
            else:
                record = write_xlsx_json(source, source_path, output_path)
            record["json_path"] = output_path.relative_to(root).as_posix()
            records.append(record)
        except Exception as exc:
            failures.append({"source_id": source["source_id"], "source_path": source["path"], "error": str(exc)})
        print(f"{index}/{len(sources)} {source['path']}")

    manifest = {
        "schema_version": "1.0.0",
        "source_file_count": len(sources),
        "converted_file_count": len(records),
        "failed_file_count": len(failures),
        "reused_file_count": reused,
        "converted_this_run": len(records) - reused,
        "table_count": sum(item["table_count"] for item in records),
        "row_count": sum(item["row_count"] for item in records),
        "cell_count": sum(item["cell_count"] for item in records),
        "records": records,
        "failures": failures,
    }
    (knowledge_dir / "tabular_data_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({key: manifest[key] for key in ("source_file_count", "converted_file_count", "reused_file_count", "converted_this_run", "failed_file_count", "table_count", "row_count", "cell_count")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
