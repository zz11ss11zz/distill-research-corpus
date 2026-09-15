from __future__ import annotations

import argparse
import copy
import csv
import gzip
import hashlib
import json
import math
import os
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

from openpyxl import load_workbook
from pypdf import PdfReader


SCHEMA_VERSION = "1.0.0"
FACT_EXTRACTION_VERSION = "1.2.0"
DOC_EXTENSIONS = {".pdf", ".md", ".txt"}
TABULAR_EXTENSIONS = {".csv", ".xlsx"}
DOI_RE = re.compile(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", re.IGNORECASE)
YEAR_RE = re.compile(r"(?<!\d)(19|20)\d{2}(?!\d)")
TEMP_RE = re.compile(r"(?<!\d)(-?\d+(?:\.\d+)?)\s*(?:°\s*)?[CK](?![A-Za-z])", re.IGNORECASE)
TABLE_TEMP_RE = re.compile(
    r"(?<!\d)(-?\d+(?:\.\d+)?(?:\s*[−–-]\s*-?\d+(?:\.\d+)?)?)\s*(?:°\s*)?[CK](?![A-Za-z])",
    re.IGNORECASE,
)
TABLE_NUMBER_RE = re.compile(
    r"(?<![A-Za-z0-9.])(?:~|∼|≈)?\d+(?:\.\d+)?(?:\s*[−–-]\s*\d+(?:\.\d+)?)?(?![A-Za-z0-9.])"
)
VALUE_UNIT_RE = re.compile(r"(?!)")
FORMULA_RE = re.compile(r"(?!)")
# A single PDF line ending is layout, not a sentence boundary. Blank lines and
# punctuation still delimit evidence; table rows are handled separately below.
SECTION_SPLIT_RE = re.compile(r"(?<=[.!?。！？;；])\s+|\n[ \t]*\n+")
TABLE_PAGE_RE = re.compile(r"(?:^|\n)\s*Table\s+(?:S?\d+|[IVX]+)\b", re.IGNORECASE)

from corpus_profile import load_profile, local_path, master_paths as configured_master_paths

KEYWORD_GROUPS = {}
PROFILE = {}
DEFAULT_CONFIG_PATH = None
ENGINE_VERSION = "2.0.0"

def configure_profile(root: Path, path: Path | None = None) -> None:
    global PROFILE, KEYWORD_GROUPS, VALUE_UNIT_RE, FORMULA_RE, FACT_EXTRACTION_VERSION
    PROFILE = load_profile(root, path)
    KEYWORD_GROUPS = {name: re.compile(pattern, re.IGNORECASE) for name,pattern in PROFILE.get("keyword_groups",{}).items()}
    VALUE_UNIT_RE = re.compile(PROFILE["value_unit_pattern"], re.IGNORECASE)
    FORMULA_RE = re.compile(PROFILE["formula_pattern"])
    FACT_EXTRACTION_VERSION = ENGINE_VERSION + ":" + PROFILE["profile_sha256"]


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def relpath(path: Path, root: Path) -> str:
    return Path(os.path.relpath(os.path.abspath(path), os.path.abspath(root))).as_posix()


def io_path(path: Path) -> Path:
    """Return a Windows extended-length path while keeping stored paths portable."""
    absolute = os.path.abspath(path)
    if os.name != "nt" or absolute.startswith("\\\\?\\"):
        return Path(absolute)
    if absolute.startswith("\\\\"):
        return Path("\\\\?\\UNC\\" + absolute[2:])
    return Path("\\\\?\\" + absolute)


def sha256_file(path: Path, chunk_size: int = 8 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest().upper()


def stable_id(prefix: str, value: str, length: int = 20) -> str:
    encoded = value.encode("utf-8", errors="replace")
    return f"{prefix}_{hashlib.sha256(encoded).hexdigest()[:length]}"


def clean_text(value: str) -> str:
    value = value.replace("\x00", " ").replace("\u00ad", "")
    value = value.encode("utf-8", errors="replace").decode("utf-8")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n[ \t]+", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (datetime,)):
        return value.isoformat()
    return str(value)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_json_array(path: Path, records: Iterable[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        stream.write("[\n")
        first = True
        for record in records:
            if not first:
                stream.write(",\n")
            json.dump(record, stream, ensure_ascii=False, indent=2)
            first = False
            count += 1
        stream.write("\n]\n")
    return count


def inventory_files(root: Path) -> list[dict[str, Any]]:
    category_dirs = {name: local_path(root,value) for name,value in PROFILE["canonical_roots"].items()}
    rows: list[dict[str, Any]] = []
    for category, base in category_dirs.items():
        if not base.exists():
            raise FileNotFoundError(f"Missing canonical category directory: {base}")
        relative_base = relpath(base, root)
        command = subprocess.run(
            ["rg", "--files", relative_base],
            cwd=root,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        if command.returncode not in (0, 1):
            raise RuntimeError(f"Unable to enumerate {relative_base}: {command.stderr.strip()}")
        enumerated = [root / Path(line.strip()) for line in command.stdout.splitlines() if line.strip()]
        for path in enumerated:
                try:
                    readable_path = io_path(path)
                    size = readable_path.stat().st_size
                    digest = sha256_file(readable_path)
                except OSError as exc:
                    rows.append(
                        {
                            "category": category,
                            "path": relpath(path, root),
                            "extension": path.suffix.lower(),
                            "size_bytes": None,
                            "sha256": None,
                            "read_error": str(exc),
                        }
                    )
                    continue
                rows.append(
                    {
                        "category": category,
                        "path": relpath(path, root),
                        "extension": path.suffix.lower(),
                        "size_bytes": size,
                        "sha256": digest,
                        "read_error": None,
                    }
                )
    return sorted(rows, key=lambda item: (item["category"], item["path"].casefold()))


def canonicalize_inventory(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    by_hash: dict[str, list[dict[str, Any]]] = defaultdict(list)
    unreadable: list[dict[str, Any]] = []
    for row in rows:
        if row["sha256"]:
            by_hash[row["sha256"]].append(row)
        else:
            unreadable.append(row)

    priority = {"paper": 0, "thesis": 1, "patent": 2, "dataset": 3}
    unique: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    for digest, group in sorted(by_hash.items()):
        group.sort(key=lambda item: (priority[item["category"]], len(item["path"]), item["path"].casefold()))
        canonical = dict(group[0])
        canonical["source_id"] = stable_id("src", digest)
        canonical["all_categories"] = sorted({item["category"] for item in group})
        canonical["alias_paths"] = [item["path"] for item in group[1:]]
        unique.append(canonical)
        for duplicate in group[1:]:
            duplicates.append(
                {
                    "sha256": digest,
                    "duplicate_path": duplicate["path"],
                    "duplicate_category": duplicate["category"],
                    "canonical_source_id": canonical["source_id"],
                    "canonical_path": canonical["path"],
                    "reason": "byte_identical_sha256",
                }
            )
    return sorted(unique, key=lambda item: item["path"].casefold()), duplicates + unreadable


def filename_metadata(path: Path, category: str) -> dict[str, Any]:
    stem = path.stem
    parts = [part.strip() for part in stem.split(" - ")]
    result: dict[str, Any] = {"title": stem, "year": None, "venue_or_institution": None}
    if category == "paper" and len(parts) >= 3:
        result.update({"venue_or_institution": parts[0], "year": int(parts[1]) if parts[1].isdigit() else None, "title": " - ".join(parts[2:])})
    elif category == "thesis" and len(parts) >= 3:
        result.update({"venue_or_institution": parts[0], "year": int(parts[1]) if parts[1].isdigit() else None, "title": " - ".join(parts[2:])})
    elif category == "patent":
        result["publication_number"] = parts[0] if parts else stem
        if len(parts) >= 2 and parts[1].isdigit():
            result["year"] = int(parts[1])
    return result


def extract_named_section(text: str, headings: list[str], stop_headings: list[str], max_chars: int = 12000) -> str | None:
    normalized = clean_text(text)
    for heading in headings:
        match = re.search(heading, normalized, re.IGNORECASE | re.MULTILINE)
        if not match:
            continue
        tail = normalized[match.end():]
        stop_positions = []
        for stop in stop_headings:
            stop_match = re.search(stop, tail, re.IGNORECASE | re.MULTILINE)
            if stop_match:
                stop_positions.append(stop_match.start())
        end = min(stop_positions) if stop_positions else min(len(tail), max_chars)
        value = clean_text(tail[: min(end, max_chars)])
        if len(value) >= 80:
            return value
    return None


def fact_candidates(pages: list[dict[str, Any]], source_id: str) -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    seen: set[tuple[Any, str]] = set()
    for page in pages:
        page_text = page["text"]
        page_chars = list(re.finditer(r"\S", page_text))
        compact_page = "".join(m.group() for m in page_chars)
        page_sha256 = hashlib.sha256(page_text.encode("utf-8")).hexdigest()
        split_rule = re.compile(r"(?<=[.!?。！？;；])\s+|\n+") if TABLE_PAGE_RE.search(page_text) else SECTION_SPLIT_RE
        candidates: list[tuple[str, str]] = [
            ("deterministic_full_text_sentence_screening", sentence)
            for sentence in split_rule.split(page_text)
        ]
        lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]
        page_groups = [name for name, pattern in KEYWORD_GROUPS.items() if pattern.search(page_text[:2500])]
        table_lines = lines if TABLE_PAGE_RE.search(page_text) else []
        for index, line in enumerate(table_lines):
            line_formulas = FORMULA_RE.findall(line)
            if not any(re.search(r"\d", formula) for formula in line_formulas):
                continue
            block = [line]
            following_lines = table_lines[index + 1:index + 9]
            for offset, following in enumerate(following_lines):
                following_formulas = FORMULA_RE.findall(following)
                if any(re.search(r"\d", formula) for formula in following_formulas):
                    break
                next_line = following_lines[offset + 1] if offset + 1 < len(following_lines) else ""
                next_formulas = FORMULA_RE.findall(next_line)
                if (
                    len(following) < 60
                    and not TABLE_NUMBER_RE.search(following)
                    and any(re.search(r"\d", formula) for formula in next_formulas)
                ):
                    break
                block.append(following)
            table_row = clean_text(" ".join(block))
            if TABLE_NUMBER_RE.search(table_row):
                candidates.append(("deterministic_pdf_table_row_screening", table_row))

        for extraction_method, candidate in candidates:
            sentence = " ".join(clean_text(candidate).split())
            if len(sentence) < 35 or len(sentence) > 1600:
                continue
            groups = [name for name, pattern in KEYWORD_GROUPS.items() if pattern.search(sentence)]
            values = [match.group(0).strip() for match in VALUE_UNIT_RE.finditer(sentence)]
            formulas = sorted(set(FORMULA_RE.findall(sentence)))
            temperatures = [match.group(0).strip() for match in TEMP_RE.finditer(sentence)]
            if extraction_method == "deterministic_pdf_table_row_screening":
                groups = sorted(set(groups + page_groups))
                values = sorted(set(values + [match.group(0).strip() for match in TABLE_NUMBER_RE.finditer(sentence)]))
                temperatures = sorted(set(match.group(0).strip() for match in TABLE_TEMP_RE.finditer(sentence)))
            if not groups and not (values and formulas):
                continue
            signature = (page["page"], sentence)
            if signature in seen:
                continue
            seen.add(signature)
            # Retain raw page offsets, without interpreting glyphs or numbers.
            needle = re.sub(r"\s+", "", sentence)
            position = compact_page.find(needle)
            locator = {}
            if position >= 0 and compact_page.find(needle, position + 1) < 0:
                start = page_chars[position].start()
                end = page_chars[position + len(needle) - 1].end()
                utf16_start = len(page_text[:start].encode("utf-16-le")) // 2
                utf16_end = len(page_text[:end].encode("utf-16-le")) // 2
                locator = {"source_span": {
                    "span_id": stable_id("span", f"{source_id}|{page['page']}|{page_sha256}|{utf16_start}|{utf16_end}"),
                    "source_id": source_id, "page": page["page"], "page_sha256": page_sha256,
                    "start": utf16_start, "end": utf16_end, "offset_unit": "utf16_code_units",
                    "raw_text_sha256": hashlib.sha256(page_text[start:end].encode("utf-8")).hexdigest(),
                }}
            facts.append(
                {
                    "fact_id": stable_id("fact", f"{source_id}|{page['page']}|{sentence}"),
                    "source_id": source_id,
                    "page": page["page"],
                    "evidence_text": sentence,
                    "themes": groups,
                    "material_formula_candidates": formulas,
                    "numeric_value_candidates": values,
                    "temperature_candidates": temperatures,
                    "extraction_method": extraction_method,
                    "verification_status": "source_located_not_human_verified",
                    "confidence": "high" if values and (groups or formulas) else "medium",
                    **locator,
                }
            )
    return facts


MACHINE_FACT_STATUS = "source_located_not_human_verified"


def is_verified_fact(fact: dict[str, Any]) -> bool:
    return fact.get("verification_status", "").startswith(("human_verified", "agent_verified"))


def load_preserved_facts(facts_path: Path, overlays: Iterable[Path] = ()) -> list[dict[str, Any]]:
    """Load the durable evidence layer even for --full; corrupt history is fatal.

    An overlay is an explicit array (or {facts: [...]}) of previously verified
    fact records, not a candidate queue. Conflicting edits need a separate,
    reviewed scientific migration rather than last-file-wins behavior.
    """
    merged: dict[str, dict[str, Any]] = {}
    for path in [facts_path, *overlays]:
        if path == facts_path and not path.exists():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        records = payload if isinstance(payload, list) else payload.get("facts")
        if not isinstance(records, list):
            raise ValueError(f"Expected fact records in {path}")
        local_ids: set[str] = set()
        for fact in records:
            if not all(fact.get(key) is not None for key in ("fact_id", "source_id", "page", "verification_status")):
                raise ValueError(f"Incomplete preserved fact in {path}")
            fid = fact["fact_id"]
            if fid in local_ids:
                raise ValueError(f"Duplicate fact ID in {path}: {fid}")
            local_ids.add(fid)
            if path != facts_path and not is_verified_fact(fact):
                raise ValueError(f"Overlay contains an unverified candidate: {fid}")
            previous = merged.get(fid)
            if previous is not None and previous != fact:
                # A machine record may be upgraded by a verified overlay only
                # when the immutable source/page/text identity is unchanged.
                identity = ("source_id", "page", "evidence_text")
                if is_verified_fact(previous) or any(previous.get(k) != fact.get(k) for k in identity):
                    raise ValueError(f"Conflicting preserved fact: {fid}")
            merged[fid] = copy.deepcopy(fact)
    return list(merged.values())


def merge_preserved_facts(
    generated: Iterable[dict[str, Any]], previous: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Retain every old ID/annotation, including fragment IDs used by decisions.

    Whitespace-only reflow reuses the old identity; new sentences remain machine
    candidates. Old fragments are retained verbatim for provenance, not upgraded
    or reinterpreted as additional independently verified measurements.
    """
    by_id = {fact["fact_id"]: fact for fact in previous}
    if len(by_id) != len(previous):
        raise ValueError("Duplicate preserved fact IDs")
    text_keys = {
        (f["source_id"], f["page"], re.sub(r"\s+", "", f.get("evidence_text", "")))
        for f in previous
    }
    result = list(previous)
    added = 0
    for fact in generated:
        key = (fact["source_id"], fact["page"], re.sub(r"\s+", "", fact.get("evidence_text", "")))
        old = by_id.get(fact["fact_id"])
        if old is not None:
            if any(old.get(k) != fact.get(k) for k in ("source_id", "page", "evidence_text")):
                raise ValueError(f"Fact ID collision: {fact['fact_id']}")
            continue
        if key in text_keys:
            continue
        result.append(fact)
        by_id[fact["fact_id"]] = fact
        text_keys.add(key)
        added += 1
    return result, {"preserved_fact_ids": len(previous), "preserved_verified_facts": sum(map(is_verified_fact, previous)), "new_machine_candidates": added}


def write_json_array_atomic(path: Path, records: Iterable[dict[str, Any]]) -> int:
    """A failed extraction must not truncate the existing evidence layer."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, prefix="kb_refresh_", suffix=".tmp", delete=False) as stream:
        pending = Path(stream.name)
    try:
        count = write_json_array(pending, records)
        os.replace(pending, path)
        return count
    finally:
        pending.unlink(missing_ok=True)


def disposition_input(output: Path, explicit: Path | None = None) -> Path | None:
    if explicit is not None:
        return explicit
    current = output / "evidence_fact_dispositions.json"
    if current.exists():
        return current
    snapshots = sorted(output.glob("evidence_fact_dispositions_*.json"))
    return snapshots[-1] if snapshots else None


def merge_fact_dispositions(facts: list[dict[str, Any]], payload: dict[str, Any]) -> dict[str, Any]:
    old_rows = payload.get("dispositions", [])
    old = {row["fact_id"]: row for row in old_rows}
    if len(old) != len(old_rows):
        raise ValueError("Duplicate disposition fact IDs")
    missing = set(old) - {fact["fact_id"] for fact in facts}
    if missing:
        raise ValueError(f"Disposition references would be lost: {sorted(missing)[:5]}")
    rows = []
    for fact in facts:
        if fact["fact_id"] in old:
            row = old[fact["fact_id"]]
            if any(row.get(k) != fact.get(k) for k in ("source_id", "page")):
                raise ValueError(f"Disposition identity conflict: {fact['fact_id']}")
            row = copy.deepcopy(row)
            # Preserve the decision independently of a verified overlay upgrade.
            row["verification_status"] = fact["verification_status"]
        else:
            row = {k: fact[k] for k in ("fact_id", "source_id", "page", "verification_status")}
            row.update({
                "formal_disposition": "human_verified_evidence_fact" if fact.get("verification_status") == "human_verified" else "verified_evidence_fact" if is_verified_fact(fact) else "machine_context_candidate_not_strict_formal_measurement",
                "reason": "retained_in_verified_knowledge_layer" if is_verified_fact(fact) else "not_selected_or_reviewed_as_a_formal_measurement",
            })
        rows.append(row)
    return {**payload, "dispositions": rows, "counts": {"total": len(rows), **Counter(r["formal_disposition"] for r in rows)}}


def migrate_cached_facts(output: Path, previous: list[dict[str, Any]], dispositions_path: Path | None,
                         destination: Path | None, apply: bool) -> dict[str, Any]:
    """Stage a bounded cached-page refresh; never publish formal data or history."""
    documents = json.loads((output / "documents.json").read_text(encoding="utf-8"))
    source_ids = {d["source_id"] for d in documents}
    if any(f["source_id"] not in source_ids for f in previous):
        raise ValueError("Preserved facts reference sources absent from cached documents")
    generated = (f for doc in documents for f in fact_candidates(doc.get("pages", []), doc["source_id"]))
    facts, counts = merge_preserved_facts(generated, previous)
    disposition_payload = json.loads(dispositions_path.read_text(encoding="utf-8")) if dispositions_path else {}
    merged_dispositions = merge_fact_dispositions(facts, disposition_payload)
    report = {"mode": "apply" if apply else "dry-run", "counts": {"before": len(previous), "after": len(facts), **counts},
              "input_sha256": {"evidence_facts": sha256_file(output / "evidence_facts.json"), "documents": sha256_file(output / "documents.json"),
                               "dispositions": sha256_file(dispositions_path) if dispositions_path else None},
              "fact_extraction_version": FACT_EXTRACTION_VERSION,
              "scientific_decisions_changed": 0, "verified_facts_changed": 0,
              "output_directory": str(destination) if destination else None}
    if apply:
        if destination is None or destination.exists():
            raise ValueError("--apply requires a new, non-existing --migration-output directory")
        destination.mkdir(parents=True)
        write_json_array_atomic(destination / "evidence_facts.json", facts)
        write_json(destination / "evidence_fact_dispositions.json", merged_dispositions)
        write_json(destination / "kb_fact_migration.json", report)
    return report


def extract_pdf(
    path: Path,
    source: dict[str, Any],
    ocr_document: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    pages: list[dict[str, Any]] = []
    errors: list[str] = []
    metadata: dict[str, Any] = {}
    try:
        reader = PdfReader(path)
        raw_metadata = reader.metadata or {}
        metadata = {
            "title": json_safe(raw_metadata.get("/Title")),
            "author": json_safe(raw_metadata.get("/Author")),
            "subject": json_safe(raw_metadata.get("/Subject")),
            "creator": json_safe(raw_metadata.get("/Creator")),
        }
        for number, page in enumerate(reader.pages, 1):
            try:
                text = clean_text(page.extract_text() or "")
            except Exception as exc:  # malformed individual pages should not lose the document
                text = ""
                errors.append(f"page {number}: {exc}")
            pages.append({"page": number, "text": text, "text_chars": len(text)})
    except Exception as exc:
        errors.append(f"document: {exc}")

    ocr_pages_used = 0
    if ocr_document:
        ocr_by_page = {int(item["page"]): clean_text(item.get("text") or "") for item in ocr_document.get("pages", [])}
        for page in pages:
            ocr_text = ocr_by_page.get(page["page"], "")
            if len(ocr_text) >= 20 and len(ocr_text) > max(20, len(page["text"]) * 2):
                page["text"] = ocr_text
                page["text_chars"] = len(ocr_text)
                page["text_source"] = "windows_ocr_zh-Hans-CN_150dpi"
                ocr_pages_used += 1
            else:
                page["text_source"] = "embedded_pdf_text"
    else:
        for page in pages:
            page["text_source"] = "embedded_pdf_text"

    all_text = "\n".join(item["text"] for item in pages)
    category = source["category"]
    title_meta = filename_metadata(path, category)
    abstract = extract_named_section(
        all_text,
        [r"^\s*abstract\s*[:—-]?", r"^\s*摘要\s*[:：]?"],
        [r"^\s*(?:keywords?|key words|关键词)\s*[:：]", r"^\s*(?:1\.?\s*)?introduction\b", r"^\s*引言\b"],
    )
    conclusion = extract_named_section(
        all_text,
        [r"^\s*(?:\d+(?:\.\d+)*\.?\s*)?(?:conclusions?|summary and conclusions?|结论与展望|结论|总结与展望)\s*$"],
        [r"^\s*(?:acknowledg|author contribution|conflict|references|参考文献)"],
    )
    facts = fact_candidates(pages, source["source_id"])
    formula_counter = Counter(formula for fact in facts for formula in fact["material_formula_candidates"])
    theme_counter = Counter(theme for fact in facts for theme in fact["themes"])
    doi_candidates = sorted({match.group(0).rstrip(".,;:)]}") for match in DOI_RE.finditer(all_text)})
    text_chars = len(all_text)
    record = {
        "source_id": source["source_id"],
        "category": category,
        "path": source["path"],
        "alias_paths": source["alias_paths"],
        "sha256": source["sha256"],
        "size_bytes": source["size_bytes"],
        "bibliographic_metadata": {**title_meta, "pdf_metadata": metadata, "doi_candidates": doi_candidates[:50]},
        "extraction": {
            "page_count": len(pages),
            "text_chars": text_chars,
            "text_coverage_status": "text_extracted" if text_chars >= max(200, len(pages) * 20) else "ocr_or_manual_review_required",
            "errors": errors,
            "ocr_pages_used": ocr_pages_used,
        },
        "distilled_sections": {"abstract_or_summary": abstract, "conclusion_or_outlook": conclusion},
        "distilled_index": {
            "themes": dict(theme_counter.most_common()),
            "material_formula_candidates": [item for item, _ in formula_counter.most_common(200)],
            "fact_count": len(facts),
        },
        "pages": pages,
        "distillation_status": (
            "machine_distilled_with_ocr_page_provenance"
            if ocr_pages_used
            else "machine_distilled_with_page_provenance"
            if text_chars
            else "unreadable_requires_ocr"
        ),
        "human_verification_status": "not_complete",
    }
    return record, facts


def extract_text_document(path: Path, source: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    errors: list[str] = []
    try:
        text = clean_text(path.read_text(encoding="utf-8", errors="replace"))
    except Exception as exc:
        text = ""
        errors.append(str(exc))
    pages = [{"page": 1, "text": text, "text_chars": len(text)}]
    facts = fact_candidates(pages, source["source_id"])
    record = {
        "source_id": source["source_id"],
        "category": source["category"],
        "path": source["path"],
        "alias_paths": source["alias_paths"],
        "sha256": source["sha256"],
        "size_bytes": source["size_bytes"],
        "bibliographic_metadata": filename_metadata(path, source["category"]),
        "extraction": {"page_count": 1, "text_chars": len(text), "text_coverage_status": "text_extracted" if text else "unreadable", "errors": errors},
        "distilled_sections": {"abstract_or_summary": None, "conclusion_or_outlook": None},
        "distilled_index": {"themes": dict(Counter(t for f in facts for t in f["themes"])), "material_formula_candidates": sorted({m for f in facts for m in f["material_formula_candidates"]}), "fact_count": len(facts)},
        "pages": pages,
        "distillation_status": "machine_distilled_with_line_level_text" if text else "unreadable",
        "human_verification_status": "not_complete",
    }
    return record, facts


def infer_scalar(value: str) -> Any:
    stripped = value.strip()
    if stripped == "":
        return None
    lowered = stripped.casefold()
    if lowered in {"true", "false"}:
        return lowered == "true"
    try:
        return int(stripped)
    except ValueError:
        try:
            number = float(stripped)
            return number if math.isfinite(number) else stripped
        except ValueError:
            return stripped


def summarize_values(values: list[Any]) -> dict[str, Any]:
    nonempty = [value for value in values if value not in (None, "")]
    numeric = [float(value) for value in nonempty if isinstance(value, (int, float)) and math.isfinite(float(value))]
    result: dict[str, Any] = {
        "nonempty_count": len(nonempty),
        "missing_count": len(values) - len(nonempty),
        "distinct_count": len({json.dumps(json_safe(value), ensure_ascii=False, sort_keys=True) for value in nonempty[:100000]}),
        "sample_values": [json_safe(value) for value in nonempty[:8]],
    }
    if numeric:
        result["numeric"] = {
            "count": len(numeric),
            "min": min(numeric),
            "max": max(numeric),
            "mean": statistics.fmean(numeric),
        }
    return result


def profile_csv(path: Path) -> dict[str, Any]:
    encodings = ["utf-8-sig", "utf-8", "gb18030", "latin-1"]
    last_error = None
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="", errors="strict") as stream:
                sample = stream.read(65536)
                stream.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(stream, dialect)
                headers = next(reader, [])
                columns = [[] for _ in headers]
                rows = 0
                malformed = 0
                for raw_row in reader:
                    rows += 1
                    if len(raw_row) != len(headers):
                        malformed += 1
                    for index in range(len(headers)):
                        value = raw_row[index] if index < len(raw_row) else ""
                        if len(columns[index]) < 100000:
                            columns[index].append(infer_scalar(value))
                return {
                    "format": "csv",
                    "encoding": encoding,
                    "delimiter": dialect.delimiter,
                    "row_count": rows,
                    "column_count": len(headers),
                    "malformed_row_count": malformed,
                    "columns": [
                        {"name": str(header), **summarize_values(columns[index])}
                        for index, header in enumerate(headers)
                    ],
                    "profiling_note": "Column statistics inspect at most the first 100000 rows; row_count covers the complete file.",
                }
        except (UnicodeError, csv.Error, OSError) as exc:
            last_error = str(exc)
    return {"format": "csv", "error": last_error or "unable_to_read"}


def profile_xlsx(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        headers_raw = next(iterator, ())
        headers = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(headers_raw)]
        columns = [[] for _ in headers]
        rows = 0
        for row in iterator:
            rows += 1
            if rows <= 100000:
                for index in range(len(headers)):
                    columns[index].append(json_safe(row[index] if index < len(row) else None))
        sheets.append(
            {
                "name": sheet.title,
                "row_count": rows,
                "column_count": len(headers),
                "columns": [{"name": header, **summarize_values(columns[index])} for index, header in enumerate(headers)],
            }
        )
    workbook.close()
    return {
        "format": "xlsx",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "profiling_note": "Column statistics inspect at most the first 100000 data rows per sheet; row_count covers the complete sheet.",
    }


def shape_of_json(value: Any, depth: int = 0) -> Any:
    if depth >= 4:
        return type(value).__name__
    if isinstance(value, dict):
        return {str(key): shape_of_json(item, depth + 1) for key, item in list(value.items())[:100]}
    if isinstance(value, list):
        return {"type": "array", "length": len(value), "item_shape": shape_of_json(value[0], depth + 1) if value else None}
    return type(value).__name__


def profile_json(path: Path) -> dict[str, Any]:
    if path.stat().st_size == 0:
        return {"format": "json", "shape": None, "content_status": "empty_file"}
    with path.open("r", encoding="utf-8-sig", errors="replace") as stream:
        payload = json.load(stream)
    return {"format": "json", "shape": shape_of_json(payload)}


def profile_gzip(path: Path) -> dict[str, Any]:
    with gzip.open(path, "rb") as stream:
        head = stream.read(1024 * 1024)
    text = head.decode("utf-8", errors="replace").lstrip()
    result: dict[str, Any] = {"format": "gzip", "decompressed_prefix_bytes_inspected": len(head)}
    if text.startswith("{") or text.startswith("["):
        try:
            with gzip.open(path, "rt", encoding="utf-8", errors="replace") as stream:
                payload = json.load(stream)
            result.update({"content_type": "json", "shape": shape_of_json(payload)})
        except Exception as exc:
            result.update({"content_type": "json_like", "parse_error": str(exc)})
    else:
        result.update({"content_type": "text_or_binary", "text_prefix": clean_text(text[:4000])})
    return result


def profile_zip(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        members = [
            {
                "path": info.filename,
                "size_bytes": info.file_size,
                "compressed_size_bytes": info.compress_size,
                "crc32": f"{info.CRC:08X}",
                "is_directory": info.is_dir(),
            }
            for info in archive.infolist()
        ]
    return {
        "format": "zip",
        "member_count": len(members),
        "uncompressed_size_bytes": sum(item["size_bytes"] for item in members),
        "members": members,
    }


def profile_code_or_text(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    return {
        "format": path.suffix.lower().lstrip(".") or "text",
        "line_count": len(lines),
        "text_chars": len(text),
        "head": clean_text("\n".join(lines[:120]))[:12000],
    }


def profile_dataset(path: Path, source: dict[str, Any]) -> dict[str, Any]:
    extension = path.suffix.lower()
    try:
        if extension == ".csv":
            profile = profile_csv(path)
        elif extension == ".xlsx":
            profile = profile_xlsx(path)
        elif extension == ".json":
            profile = profile_json(path)
        elif extension == ".gz":
            profile = profile_gzip(path)
        elif extension == ".zip":
            profile = profile_zip(path)
        elif extension in {".py", ".md", ".txt"}:
            profile = profile_code_or_text(path)
        elif extension == ".pdf":
            profile = {"format": "pdf", "handled_as_document": True}
        else:
            profile = {"format": extension.lstrip(".") or "unknown", "profiling_status": "binary_registered_only"}
        status = "profiled"
        error = None
    except Exception as exc:
        profile = {"format": extension.lstrip(".") or "unknown"}
        status = "profile_failed"
        error = str(exc)
    return {
        "source_id": source["source_id"],
        "category": "dataset",
        "path": source["path"],
        "alias_paths": source["alias_paths"],
        "sha256": source["sha256"],
        "size_bytes": source["size_bytes"],
        "profile": profile,
        "distillation_status": status,
        "error": error,
        "human_verification_status": "not_complete",
    }


def master_collection_profile(root: Path, path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        payload = json.load(stream)
    return {
        "collection_id": stable_id("collection", relpath(path, root)),
        "path": relpath(path, root),
        "sha256": sha256_file(path),
        "size_bytes": path.stat().st_size,
        "schema_version": payload.get("schema_version"),
        "database_type": payload.get("database_type"),
        "statistics": payload.get("statistics"),
        "record_counts": {
            "records": len(payload.get("records", [])),
            "composition_only_records": len(payload.get("composition_only_records", [])),
            "screened_literature_records": len(payload.get("screened_literature_extracted_data", {}).get("records", [])),
        },
        "integration_mode": "authoritative_existing_json_referenced_without_copying",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Distill a configured research corpus into deduplicated JSON.")
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, default=Path("outputs/json/knowledge_distillation"))
    parser.add_argument("--inventory-only", action="store_true")
    parser.add_argument("--full", action="store_true", help="Ignore prior JSON records and re-extract every unique source.")
    parser.add_argument("--refresh-facts", action="store_true", help="Rebuild facts from cached page text without re-reading unchanged PDFs.")
    parser.add_argument("--verified-overlay", type=Path, action="append", default=[], help="Explicit verified fact array (or {facts: [...]}) to preserve/restore; repeatable.")
    parser.add_argument("--dispositions", type=Path, help="Existing disposition snapshot to preserve; defaults to current/latest local snapshot.")
    parser.add_argument("--migrate-facts", action="store_true", help="Only refresh cached-page candidates and merge existing IDs/decisions; dry-run by default.")
    parser.add_argument("--migration-output", type=Path, help="New staging directory for --migrate-facts --apply; never overwritten.")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--apply", action="store_true")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    args = parser.parse_args()
    root = args.root.resolve()
    configure_profile(root, args.config)
    output = (root / args.output).resolve() if not args.output.is_absolute() else args.output.resolve()
    generated_at = utc_now()
    if (args.dry_run or args.apply or args.migration_output) and not args.migrate_facts:
        parser.error("--dry-run/--apply/--migration-output require --migrate-facts")
    previous_facts = load_preserved_facts(output / "evidence_facts.json", [root / p for p in args.verified_overlay])
    dispositions_path = disposition_input(output, root / args.dispositions if args.dispositions else None)
    if args.migrate_facts:
        if args.full or args.inventory_only:
            parser.error("--migrate-facts cannot be combined with --full/--inventory-only")
        report = migrate_cached_facts(output, previous_facts, dispositions_path,
                                      root / args.migration_output if args.migration_output else None, args.apply)
        print(json.dumps(report, ensure_ascii=False))
        return

    ocr_path = output / "ocr_pages.json"
    ocr_by_hash: dict[str, dict[str, Any]] = {}
    if ocr_path.exists():
        ocr_payload = json.loads(ocr_path.read_text(encoding="utf-8"))
        ocr_by_hash = {item["sha256"]: item for item in ocr_payload.get("documents", [])}

    inventory = inventory_files(root)
    unique, duplicates = canonicalize_inventory(inventory)
    document_sources = {s["source_id"] for s in unique if s["extension"] in DOC_EXTENSIONS}
    if not args.inventory_only and any(f["source_id"] not in document_sources for f in previous_facts):
        raise ValueError("A preserved evidence source is absent from inventory; restore/explicitly migrate it before rebuilding")
    write_json(output / "source_inventory.json", {"schema_version": SCHEMA_VERSION, "generated_at": generated_at, "files": inventory})
    write_json(output / "deduplication.json", {"schema_version": SCHEMA_VERSION, "generated_at": generated_at, "duplicate_or_unreadable_entries": duplicates})
    if args.inventory_only:
        print(json.dumps({"files": len(inventory), "unique": len(unique), "duplicate_or_unreadable": len(duplicates)}, ensure_ascii=False))
        return

    documents_path = output / "documents.json"
    facts_path = output / "evidence_facts.json"
    datasets_path = output / "datasets.json"
    cached_documents: dict[str, dict[str, Any]] = {}
    cached_facts: dict[str, list[dict[str, Any]]] = defaultdict(list)
    cached_datasets: dict[str, dict[str, Any]] = {}
    cached_fact_version: str | None = None
    for item in previous_facts:
        cached_facts[item["source_id"]].append(item)
    if not args.full:
        try:
            cached_manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
            cached_fact_version = cached_manifest.get("quality_model", {}).get("fact_extraction_version")
            cached_documents = {
                item["sha256"]: item for item in json.loads(documents_path.read_text(encoding="utf-8"))
            }
            cached_datasets = {
                item["sha256"]: item for item in json.loads(datasets_path.read_text(encoding="utf-8"))
            }
        except (FileNotFoundError, json.JSONDecodeError, KeyError, TypeError):
            cached_documents = {}
            cached_datasets = {}
    refresh_facts = args.refresh_facts or cached_fact_version != FACT_EXTRACTION_VERSION
    document_count = 0
    fact_count = 0
    dataset_count = 0
    unreadable_documents = 0
    failed_datasets = 0
    category_counts = Counter()
    document_reused = 0
    document_extracted = 0
    dataset_reused = 0
    dataset_profiled = 0
    fact_sources_refreshed = 0
    fact_sources_reused = 0
    all_merged_facts: list[dict[str, Any]] = []
    preservation_counts: Counter = Counter()
    pending_facts: Path | None = None

    def document_records() -> Iterator[dict[str, Any]]:
        nonlocal document_count, fact_count, unreadable_documents, document_reused, document_extracted
        nonlocal fact_sources_refreshed, fact_sources_reused
        nonlocal pending_facts
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n", dir=output,
                                         prefix="kb_refresh_", suffix=".tmp", delete=False) as fact_stream:
            pending_facts = Path(fact_stream.name)
            fact_stream.write("[\n")
            first_fact = True
            for source in unique:
                extension = source["extension"]
                if extension not in DOC_EXTENSIONS or source["category"] == "dataset" and extension not in {".pdf", ".md", ".txt"}:
                    continue
                cached = cached_documents.get(source["sha256"])
                needs_ocr_refresh = bool(
                    cached
                    and source["sha256"] in ocr_by_hash
                    and cached.get("extraction", {}).get("text_coverage_status", "").startswith("ocr")
                )
                if cached and not needs_ocr_refresh:
                    record = dict(cached)
                    record.update({
                        "source_id": source["source_id"],
                        "category": source["category"],
                        "path": source["path"],
                        "alias_paths": source["alias_paths"],
                        "sha256": source["sha256"],
                        "size_bytes": source["size_bytes"],
                    })
                    if refresh_facts:
                        facts = fact_candidates(record.get("pages", []), source["source_id"])
                        fact_sources_refreshed += 1
                    else:
                        facts = cached_facts.get(source["source_id"], [])
                        fact_sources_reused += 1
                    document_reused += 1
                else:
                    path = io_path(root / Path(source["path"]))
                    if extension == ".pdf":
                        record, facts = extract_pdf(path, source, ocr_by_hash.get(source["sha256"]))
                    else:
                        record, facts = extract_text_document(path, source)
                    document_extracted += 1
                facts, preserved = merge_preserved_facts(facts, cached_facts.get(source["source_id"], []))
                preservation_counts.update(preserved)
                all_merged_facts.extend(facts)
                document_count += 1
                category_counts[record["category"]] += 1
                if record["distillation_status"].startswith("unreadable") or record["extraction"]["text_coverage_status"].startswith("ocr"):
                    unreadable_documents += 1
                for fact in facts:
                    if not first_fact:
                        fact_stream.write(",\n")
                    json.dump(fact, fact_stream, ensure_ascii=False, indent=2)
                    first_fact = False
                    fact_count += 1
                yield record
            fact_stream.write("\n]\n")
        os.replace(pending_facts, facts_path)

    # Validate disposition identities before replacing either evidence artifact.
    old_dispositions = json.loads(dispositions_path.read_text(encoding="utf-8")) if dispositions_path else {}
    merge_fact_dispositions(previous_facts, old_dispositions)
    try:
        write_json_array_atomic(documents_path, document_records())
    finally:
        if pending_facts is not None:
            pending_facts.unlink(missing_ok=True)
    if dispositions_path:
        write_json(output / "evidence_fact_dispositions.json", merge_fact_dispositions(all_merged_facts, old_dispositions))

    def dataset_records() -> Iterator[dict[str, Any]]:
        nonlocal dataset_count, failed_datasets, dataset_reused, dataset_profiled
        for source in unique:
            if source["category"] != "dataset":
                continue
            cached = cached_datasets.get(source["sha256"])
            if cached:
                record = dict(cached)
                record.update({
                    "source_id": source["source_id"],
                    "path": source["path"],
                    "alias_paths": source["alias_paths"],
                    "sha256": source["sha256"],
                    "size_bytes": source["size_bytes"],
                })
                dataset_reused += 1
            else:
                path = io_path(root / Path(source["path"]))
                record = profile_dataset(path, source)
                dataset_profiled += 1
            dataset_count += 1
            category_counts["dataset"] += 1
            if record["distillation_status"] != "profiled":
                failed_datasets += 1
            yield record

    write_json_array(datasets_path, dataset_records())

    master_paths = configured_master_paths(PROFILE, root, output)
    master_collections = [master_collection_profile(root, path) for path in master_paths if path.exists()]
    write_json(output / "master_collections.json", {"schema_version": SCHEMA_VERSION, "generated_at": generated_at, "collections": master_collections})

    tabular_manifest_path = output / "tabular_data_manifest.json"
    tabular_statistics: dict[str, Any] = {}
    if tabular_manifest_path.exists():
        tabular_manifest = json.loads(tabular_manifest_path.read_text(encoding="utf-8"))
        tabular_statistics = {
            "tabular_json_files": tabular_manifest.get("converted_file_count", 0),
            "tabular_tables": tabular_manifest.get("table_count", 0),
            "tabular_rows": tabular_manifest.get("row_count", 0),
            "tabular_cells": tabular_manifest.get("cell_count", 0),
            "tabular_conversion_failures": tabular_manifest.get("failed_file_count", 0),
        }

    source_hashes = {item["sha256"] for item in inventory if item["sha256"]}
    unique_source_category_counts = Counter(item["category"] for item in unique)
    manifest = {
        "$schema": "knowledge_base.schema.json",
        "schema_version": SCHEMA_VERSION,
        "generated_at": generated_at,
        "scope": {
            "root": root.as_posix(),
            "canonical_categories": list(PROFILE["canonical_roots"]),
            "material_system": PROFILE.get("material_system"),
            "profile_sha256": PROFILE["profile_sha256"],
            "evidence_rules": PROFILE["evidence_rules"],
            "policy": "Every canonical file is inventoried; byte-identical files are represented once with alias paths.",
        },
        "statistics": {
            "source_file_entries": len(inventory),
            "source_unique_sha256": len(source_hashes),
            "unique_source_records": len(unique),
            "byte_identical_aliases": sum(1 for item in duplicates if item.get("reason") == "byte_identical_sha256"),
            "unreadable_inventory_entries": sum(1 for item in inventory if item["sha256"] is None),
            "document_records": document_count,
            "document_records_reused": document_reused,
            "document_records_extracted": document_extracted,
            "evidence_fact_candidates": fact_count,
            "fact_sources_refreshed": fact_sources_refreshed,
            "fact_sources_reused": fact_sources_reused,
            **preservation_counts,
            "dataset_records": dataset_count,
            "dataset_records_reused": dataset_reused,
            "dataset_records_profiled": dataset_profiled,
            "documents_requiring_ocr_or_manual_review": unreadable_documents,
            "dataset_profiles_failed": failed_datasets,
            "unique_sources_by_category": dict(unique_source_category_counts),
            "output_records_by_category": dict(category_counts),
            "master_collection_count": len(master_collections),
            **tabular_statistics,
        },
        "artifacts": {
            "source_inventory": "source_inventory.json",
            "deduplication": "deduplication.json",
            "documents": "documents.json",
            "evidence_facts": "evidence_facts.json",
            "datasets": "datasets.json",
            "tabular_data": "tabular_data_manifest.json",
            "master_collections": "master_collections.json",
            **({"evidence_fact_dispositions": "evidence_fact_dispositions.json"} if dispositions_path else {}),
        },
        "quality_model": {
            "document_text": "Complete page text is retained when extractable.",
            "fact_candidates": "Page-located deterministic extraction; not promoted to human-verified measurement facts.",
            "fact_extraction_version": FACT_EXTRACTION_VERSION,
            "fact_refresh_preservation": "Existing IDs, verified overlays and dispositions are retained; new reflowed sentences remain machine candidates.",
            "dataset_profiles": "Complete file/member/row counts with bounded column statistics; every unique CSV/XLSX cell is also converted to JSON under tabular/.",
            "human_verification": "not_complete_for_all_records",
        },
    }
    write_json(output / "manifest.json", manifest)
    print(json.dumps(manifest["statistics"], ensure_ascii=False))


if __name__ == "__main__":
    main()
